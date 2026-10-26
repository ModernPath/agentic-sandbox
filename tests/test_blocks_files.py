import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from sandbox_blocks.files import MAX_FILE_BYTES, FileError, read_document

PNG_1x1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d4944415478da63f8cfc0f01f0005000201"
    "a7f2a2b40000000049454e44ae426082")


def xlsx_bytes(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx_bytes(paragraphs, table=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx_bytes(slides):
    prs = Presentation()
    for title, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pdf_bytes(pages):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def test_csv_is_parsed_into_rows_with_header():  # F-1
    result = read_document(b"name,role\nAda,CTO\nAlan,VP\n", "people.csv")
    assert result.kind == "csv"
    assert result.tables == {"csv": [["name", "role"], ["Ada", "CTO"], ["Alan", "VP"]]}
    assert result.truncated is False


def test_csv_rows_are_capped_and_flagged(  ):  # F-2
    data = "\n".join(f"{i},x" for i in range(1000)).encode()
    result = read_document(data, "big.csv", max_rows=10)
    assert len(result.tables["csv"]) == 10 and result.truncated is True


def test_json_is_pretty_printed_text():  # F-3
    result = read_document(b'{"b": 1, "a": [1, 2]}', "data.json")
    assert result.kind == "json"
    assert json.loads(result.text) == {"b": 1, "a": [1, 2]}


def test_text_and_markdown_are_decoded():  # F-4
    assert read_document("Hei maailma\n".encode(), "notes.txt").text == "Hei maailma\n"
    assert read_document(b"# Title\n", "readme.md").kind == "text"


def test_xlsx_reads_every_sheet_as_rows():  # F-5
    data = xlsx_bytes({"Targets": [["objective", "value"], ["pipeline", 10]], "Wins": [["customer"], ["Acme"]]})
    result = read_document(data, "plan.xlsx")
    assert result.kind == "xlsx"
    assert result.tables["Targets"] == [["objective", "value"], ["pipeline", "10"]]
    assert result.tables["Wins"] == [["customer"], ["Acme"]]


def test_docx_reads_paragraphs_and_tables():  # F-6
    data = docx_bytes(["Quarterly review", "Pipeline is up."], table=[["Region", "Deals"], ["EMEIA", "12"]])
    result = read_document(data, "review.docx")
    assert result.kind == "docx"
    assert "Quarterly review" in result.text and "Pipeline is up." in result.text
    assert "EMEIA" in result.text and "12" in result.text


def test_pptx_reads_slide_text_and_notes():  # F-7
    data = pptx_bytes([("Market", "Competitor X launched")])
    result = read_document(data, "deck.pptx")
    assert result.kind == "pptx" and result.pages == 1
    assert "Market" in result.text and "Competitor X launched" in result.text


def test_pdf_reports_page_count():  # F-8
    result = read_document(pdf_bytes(3), "report.pdf")
    assert result.kind == "pdf" and result.pages == 3


def test_png_is_an_image_for_vision_not_text():  # F-9
    result = read_document(PNG_1x1, "chart.png")
    assert result.kind == "image" and result.mime == "image/png"
    assert result.text == "" and result.image_bytes == PNG_1x1


def test_file_over_10_mb_is_rejected_before_parsing():  # F-10
    assert MAX_FILE_BYTES == 10 * 1024 * 1024
    with pytest.raises(FileError, match="10 MB"):
        read_document(b"a" * (MAX_FILE_BYTES + 1), "big.txt")


@pytest.mark.parametrize("name", ["tool.exe", "page.html", "archive.zip", "noextension"])
def test_unsupported_types_are_rejected(name):  # F-11
    with pytest.raises(FileError, match="not supported"):
        read_document(b"data", name)


def test_extension_must_match_content():  # F-11b
    with pytest.raises(FileError):
        read_document(b"not a png at all", "fake.png")


def test_office_zip_with_huge_declared_size_is_rejected():  # F-12
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * (60 * 1024 * 1024))
    with pytest.raises(FileError, match="too large when unpacked"):
        read_document(buf.getvalue(), "bomb.docx")


def test_corrupt_pdf_is_a_file_error_not_a_crash():  # F-13
    with pytest.raises(FileError):
        read_document(b"%PDF-1.4 truncated", "broken.pdf")
